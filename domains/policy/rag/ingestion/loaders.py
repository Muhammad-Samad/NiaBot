"""
Document loaders. Each loader returns plain text with one logical line per
line, so the cleaner and chunker can work the same way for every format.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class LoadedDocument:
    source: str          # file name, stored in chunk metadata
    text: str
    page_count: int | None = None


def _load_pdf(path: Path) -> LoadedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    # Pages are joined with a single newline: sections and even lists routinely
    # continue across pages, and the PDF already marks real paragraph breaks
    # with blank lines.
    pages = [(page.extract_text() or "") for page in reader.pages]
    return LoadedDocument(source=path.name, text="\n".join(pages), page_count=len(pages))


def _docx_list_level(para) -> int | None:
    """List indent level of a Word paragraph, or None if it is not a list item.
    Numbering can be set on the paragraph itself or inherited from its style."""
    num_pr = para._p.pPr.numPr if para._p.pPr is not None else None
    style = para.style
    while num_pr is None and style is not None:
        ppr = style.element.pPr
        num_pr = ppr.numPr if ppr is not None else None
        style = style.base_style
    if num_pr is None:
        return None
    return num_pr.ilvl.val if num_pr.ilvl is not None else 0


def _load_docx(path: Path) -> LoadedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required to load .docx files: pip install python-docx") from exc

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines: list[str] = []
    # Walk the body in document order so tables (history, delivery-charge
    # summary) stay inside the section they belong to.
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(element, document)
            text = para.text
            # Word list items carry no bullet glyph in their text; emit the
            # same "•" / "o" markers a PDF export has so the cleaner and
            # chunker treat both formats identically.
            level = _docx_list_level(para)
            if level is not None and text.strip():
                text = f"{'o' if level else '•'} {text.strip()}"
            lines.append(text)
        elif tag == "tbl":
            table = Table(element, document)
            lines.append("")
            for row in table.rows:
                cells = []
                for cell in row.cells:  # merged cells repeat; keep one copy
                    value = " ".join(cell.text.split())
                    if not cells or cells[-1] != value:
                        cells.append(value)
                lines.append(" | ".join(cells))
            lines.append("")
    return LoadedDocument(source=path.name, text="\n".join(lines))


def _load_text(path: Path) -> LoadedDocument:
    return LoadedDocument(source=path.name, text=path.read_text(encoding="utf-8"))


_LOADERS = {".pdf": _load_pdf, ".docx": _load_docx, ".txt": _load_text, ".md": _load_text}


def load_document(path: Path) -> LoadedDocument:
    ext = path.suffix.lower()
    if ext not in _LOADERS:
        raise ValueError(f"Unsupported file type: {path.name}")
    doc = _LOADERS[ext](path)
    logger.info("Loaded %s (%d chars)", path.name, len(doc.text))
    return doc


def load_directory(directory: Path) -> list[LoadedDocument]:
    if not directory.exists():
        raise FileNotFoundError(f"Documents folder not found: {directory}")
    files = sorted(
        p for p in directory.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS and not p.name.startswith("~$")
    )
    if not files:
        raise FileNotFoundError(f"No supported documents ({', '.join(SUPPORTED_EXTENSIONS)}) in {directory}")
    return [load_document(p) for p in files]
